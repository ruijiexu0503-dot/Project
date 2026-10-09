#!/usr/bin/env python3
"""Shadow audit of native PPTX objects, geometry, z-order and overlap.

This is not a PowerPoint renderer.  The HTML is an approximate object-layout
view; exact geometry and ordering come from OOXML through python-pptx.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import io
import json
from collections import Counter
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


STYLE = """<!doctype html><meta charset="utf-8"><title>PPTX native-object audit</title>
<style>body{font:15px system-ui;background:#edf1f5;margin:18px}header,section{background:#fff;padding:16px;margin:14px auto;max-width:1280px}.slide{position:relative;background:white;box-shadow:0 1px 8px #777;width:960px;height:720px;overflow:hidden}.obj{position:absolute;box-sizing:border-box;overflow:hidden}.picture{outline:3px solid #f28c00}.text{outline:2px solid #17a768;background:rgba(255,255,255,.04);white-space:pre-wrap}.other{outline:2px dashed #7759d6}.tag{position:absolute;top:0;left:0;background:rgba(0,0,0,.72);color:white;font:11px monospace;padding:2px 4px;z-index:9999}.index td,.index th{padding:7px;border-bottom:1px solid #ddd;text-align:left}.index{border-collapse:collapse}pre{white-space:pre-wrap;max-height:520px;overflow:auto}.warn{color:#a44}</style>"""


def bbox(shape, width, height):
    return [shape.left / width, shape.top / height,
            (shape.left + shape.width) / width, (shape.top + shape.height) / height]


def intersection(first, second):
    x1, y1 = max(first[0], second[0]), max(first[1], second[1])
    x2, y2 = min(first[2], second[2]), min(first[3], second[3])
    return max(0, x2 - x1) * max(0, y2 - y1)


def area(box):
    return max(0, box[2] - box[0]) * max(0, box[3] - box[1])


def shape_kind(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return 'picture'
    if getattr(shape, 'has_text_frame', False) and shape.text.strip():
        return 'text'
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return 'group'
    return 'other'


def text_style(shape):
    sizes = []
    for paragraph in getattr(shape, 'text_frame', []).paragraphs if getattr(shape, 'has_text_frame', False) else []:
        for run in paragraph.runs:
            if run.font.size:
                sizes.append(run.font.size.pt)
    return max(9, min(34, sum(sizes) / len(sizes))) if sizes else 16


def group_children(shape, slide_width, slide_height):
    if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
        return []
    rows = []
    # Child positions are in group coordinates.  Map the group's child extent
    # into its slide placement without flattening the individual objects.
    xfrm = shape._element.grpSpPr.xfrm
    ch_x, ch_y = int(xfrm.chOff.x), int(xfrm.chOff.y)
    ch_w, ch_h = max(1, int(xfrm.chExt.cx)), max(1, int(xfrm.chExt.cy))
    for z, child in enumerate(shape.shapes):
        left = shape.left + (child.left - ch_x) * shape.width / ch_w
        top = shape.top + (child.top - ch_y) * shape.height / ch_h
        width = child.width * shape.width / ch_w
        height = child.height * shape.height / ch_h
        rows.append({'z_order_within_group': z, 'name': child.name,
                     'shape_type': str(child.shape_type), 'kind': shape_kind(child),
                     'bbox_normalized': [left / slide_width, top / slide_height,
                                         (left + width) / slide_width, (top + height) / slide_height],
                     'text_preview': child.text.strip()[:240] if getattr(child, 'has_text_frame', False) else ''})
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--pptx', required=True, type=Path)
    ap.add_argument('--output', required=True, type=Path)
    args = ap.parse_args(); args.output.mkdir(parents=True, exist_ok=True)
    data = args.pptx.read_bytes(); presentation = Presentation(io.BytesIO(data))
    sw, sh = presentation.slide_width, presentation.slide_height
    all_shapes, all_pictures, slide_summaries, index_rows = [], [], [], []
    totals = Counter()
    for slide_number, slide in enumerate(presentation.slides, 1):
        slide_dir = args.output / f'slide_{slide_number:04d}'; slide_dir.mkdir(exist_ok=True)
        rows = []
        for z_order, shape in enumerate(slide.shapes):
            kind = shape_kind(shape); box = bbox(shape, sw, sh)
            row = {'slide': slide_number, 'object_id': f'slide_{slide_number:04d}_object_{z_order + 1:03d}',
                   'z_order': z_order, 'name': shape.name, 'kind': kind,
                   'shape_type': str(shape.shape_type), 'bbox_normalized': box,
                   'rotation_degrees': float(shape.rotation or 0),
                   'text': shape.text if getattr(shape, 'has_text_frame', False) else '',
                   'group_children': group_children(shape, sw, sh)}
            if kind == 'picture':
                image = shape.image; ext = image.ext or 'bin'
                filename = f'{row["object_id"]}.{ext}'
                (slide_dir / filename).write_bytes(image.blob)
                with Image.open(io.BytesIO(image.blob)) as opened:
                    bands = opened.getbands()
                    has_alpha = 'A' in bands or ('transparency' in opened.info)
                    alpha_nonopaque_fraction = 0.0
                    if has_alpha:
                        rgba = opened.convert('RGBA')
                        alpha = rgba.getchannel('A')
                        histogram = alpha.histogram()
                        alpha_nonopaque_fraction = 1 - histogram[255] / max(1, sum(histogram))
                    source_size = list(opened.size); source_mode = opened.mode
                row.update({'asset_path': f'slide_{slide_number:04d}/{filename}',
                            'source_filename': image.filename, 'source_size': source_size,
                            'source_mode': source_mode, 'source_has_alpha': has_alpha,
                            'source_alpha_nonopaque_fraction': alpha_nonopaque_fraction,
                            'crop': {'left': float(shape.crop_left), 'top': float(shape.crop_top),
                                     'right': float(shape.crop_right), 'bottom': float(shape.crop_bottom)},
                            'layer_geometry_status': 'native_pptx_picture_shape',
                            'visible_mask_status': 'unavailable_without_faithful_renderer',
                            'hidden_extent_recoverable_from_blob': True})
                all_pictures.append(row)
            rows.append(row); all_shapes.append(row); totals[kind] += 1
        overlaps = []
        for i, first in enumerate(rows):
            for second in rows[i + 1:]:
                overlap = intersection(first['bbox_normalized'], second['bbox_normalized'])
                if overlap <= 0:
                    continue
                item = {'lower_object_id': first['object_id'], 'upper_object_id': second['object_id'],
                        'lower_kind': first['kind'], 'upper_kind': second['kind'],
                        'intersection_slide_fraction': overlap,
                        'lower_coverage': overlap / max(area(first['bbox_normalized']), 1e-12),
                        'upper_coverage': overlap / max(area(second['bbox_normalized']), 1e-12)}
                overlaps.append(item)
        picture_overlap = [o for o in overlaps if 'picture' in (o['lower_kind'], o['upper_kind'])]
        summary = {'slide': slide_number, 'objects': len(rows),
                   'pictures': sum(r['kind'] == 'picture' for r in rows),
                   'text_objects': sum(r['kind'] == 'text' for r in rows),
                   'groups': sum(r['kind'] == 'group' for r in rows),
                   'overlap_pairs': len(overlaps), 'picture_overlap_pairs': len(picture_overlap)}
        slide_summaries.append(summary)
        objects_html = []
        for row in rows:
            x1, y1, x2, y2 = row['bbox_normalized']; style = (f'left:{x1*960:.2f}px;top:{y1*720:.2f}px;'
                                                               f'width:{(x2-x1)*960:.2f}px;height:{(y2-y1)*720:.2f}px;'
                                                               f'z-index:{row["z_order"] + 1};transform:rotate({row["rotation_degrees"]}deg)')
            tag = f'<span class="tag">{row["object_id"].split("_")[-1]} z={row["z_order"]} {html.escape(row["kind"])}</span>'
            if row['kind'] == 'picture':
                content = f'<img src="{Path(row["asset_path"]).name}" style="width:100%;height:100%;object-fit:fill">'
            elif row['kind'] == 'text':
                content = f'<div style="padding:8px;font-size:{text_style(slide.shapes[row["z_order"]]):.1f}px">{html.escape(row["text"])}</div>'
            else:
                content = ''
            objects_html.append(f'<div class="obj {row["kind"] if row["kind"] in ("picture","text") else "other"}" style="{style}">{content}{tag}</div>')
        page = STYLE + f'<header><h1>Slide {slide_number}</h1><p class="warn">Approximate HTML object layout, not a faithful PowerPoint render. Geometry and z-order are native OOXML.</p></header>'
        page += '<section><div class="slide">' + ''.join(objects_html) + '</div></section>'
        page += '<section><h2>Native objects and overlaps</h2><pre>' + html.escape(json.dumps({'summary': summary, 'objects': rows, 'overlaps': overlaps}, indent=2)) + '</pre></section>'
        (slide_dir / 'report.html').write_text(page)
        index_rows.append(f'<tr><td>{slide_number}</td><td>{summary["objects"]}</td><td>{summary["pictures"]}</td><td>{summary["picture_overlap_pairs"]}</td><td><a href="slide_{slide_number:04d}/report.html">open</a></td></tr>')
    result = {'experiment': 'pptx-native-object-adapter-shadow-v1',
              'input': str(args.pptx.resolve()), 'input_sha256': hashlib.sha256(data).hexdigest(),
              'slides': len(presentation.slides), 'slide_size_emu': [sw, sh], 'totals': dict(totals),
              'pictures': len(all_pictures),
              'slides_with_picture_overlap': sum(s['picture_overlap_pairs'] > 0 for s in slide_summaries),
              'picture_overlap_pairs': sum(s['picture_overlap_pairs'] for s in slide_summaries),
              'native_capabilities': {'object_geometry': True, 'z_order': True, 'group_members': True,
                                      'original_picture_blob': True, 'picture_crop': True,
                                      'faithful_slide_render': False, 'visible_mask': False},
              'production_modified': False,
              'limitations': ['HTML preview is approximate and cannot validate PowerPoint rendering effects.',
                              'Animations, transitions, SmartArt rendering, theme fidelity and effects are not resolved.',
                              'Visible masks need a faithful renderer; they are not fabricated from bounding boxes.'],
              'slide_summaries': slide_summaries}
    (args.output / 'summary.json').write_text(json.dumps(result, indent=2))
    (args.output / 'native_objects.jsonl').write_text(''.join(json.dumps(r) + '\n' for r in all_shapes))
    (args.output / 'picture_assets.jsonl').write_text(''.join(json.dumps(r) + '\n' for r in all_pictures))
    table = '<table class="index"><tr><th>Slide</th><th>Objects</th><th>Pictures</th><th>Picture overlaps</th><th>Report</th></tr>' + ''.join(index_rows) + '</table>'
    index = STYLE + '<header><h1>PPTX native-object audit</h1><p>Orange: picture; green: text; purple dashed: other/group. Geometry and z-order come from OOXML; visual styling is approximate.</p><pre>' + html.escape(json.dumps({k: result[k] for k in ('slides','totals','pictures','slides_with_picture_overlap','picture_overlap_pairs','native_capabilities','limitations')}, indent=2)) + '</pre></header><section>' + table + '</section>'
    (args.output / 'report.html').write_text(index)
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
