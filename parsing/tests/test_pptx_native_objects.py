"""Small native PPTX adapter integration test."""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


class PptxNativeObjects(unittest.TestCase):
    def test_picture_and_upper_text_keep_native_z_order(self):
        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temporary:
            temporary = Path(temporary); image = temporary / 'image.png'
            Image.new('RGBA', (100, 100), (255, 0, 0, 180)).save(image)
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(4), Inches(3))
            box = slide.shapes.add_textbox(Inches(2), Inches(2), Inches(3), Inches(1)); box.text = 'over image'
            pptx = temporary / 'fixture.pptx'; output = temporary / 'audit'; deck.save(pptx)
            subprocess.run([sys.executable, str(root / 'src' / 'audit_pptx_native_objects.py'),
                            '--pptx', str(pptx), '--output', str(output)], check=True,
                           stdout=subprocess.DEVNULL)
            summary = json.loads((output / 'summary.json').read_text())
            pictures = [json.loads(x) for x in (output / 'picture_assets.jsonl').read_text().splitlines()]
        self.assertEqual(summary['pictures'], 1)
        self.assertEqual(summary['picture_overlap_pairs'], 1)
        self.assertTrue(pictures[0]['source_has_alpha'])
        self.assertEqual(pictures[0]['visible_mask_status'], 'unavailable_without_faithful_renderer')


if __name__ == '__main__':
    unittest.main()
