#!/usr/bin/env python3
"""Isolated EvidenceNet parsing -> official PosterGen agents -> editable poster.

Uses the existing local Transformers client; no cloud credentials or Marker.
Original prompts, layout, styling, and renderer come from pinned PosterGen.
"""
from __future__ import annotations
import argparse, base64, dataclasses, hashlib, json, os, re, subprocess, sys, time, types
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
VENDOR = ROOT / 'external/PosterGen'
sys.path[:0] = [str(HERE/'deps'), str(ROOT/'EvidenceNet'), str(VENDOR)]

def write(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False, default=lambda x: dataclasses.asdict(x) if dataclasses.is_dataclass(x) else str(x)))

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()

def prepare(out):
    import fitz
    source = ROOT/'parsing/output/hybrid_deepseek_layout_split_render/enriched_json/detr'
    pdf_path = ROOT/'data/processed/detr.pdf'
    pdf = fitz.open(pdf_path)
    assets = out/'assets'; assets.mkdir(parents=True, exist_ok=True)
    content = out/'content'; content.mkdir(parents=True, exist_ok=True)
    raw, records, figures, tables = [], [], {}, {}
    # A bounded pilot: the main paper, excluding references and appendix.
    # Explicit five-visual selection avoids ambiguous multi-figure caption matching.
    selection = {2: ('figure', '1'), 7: ('figure', '2'), 11: ('figure','3'),
                 9: ('table','1'), 16: ('table','5')}
    for num in range(1,18):
        path = source/f'page_{num:04d}.json'; data = json.loads(path.read_text())
        blocks = sorted(data['aligned_blocks'], key=lambda b:b.get('final_order',b.get('deepseek_order',0)))
        text = '\n\n'.join(b.get('display_markdown') or b.get('markdown') or b.get('text','') for b in blocks)
        if num == 17: text = re.split(r'(?im)^\s*(?:#+\s*)?(?:\*\*)?References',text)[0]
        raw.append(f'<!-- source: {path.relative_to(ROOT)}, PDF page {num} -->\n{text}')
        records.append({'source_json':str(path.relative_to(ROOT)),'sha256':sha(path),'pdf_page':num,
                        'review_items':len(data.get('page_review_items',[]))})
        if num not in selection: continue
        kind, asset_id = selection[num]
        regions = [v for v in data['visual_regions'] if v['label'] in (['image','chart'] if kind=='figure' else ['table'])]
        captions = [v for v in data['visual_regions'] if v['label'] == ('figure_title' if kind=='figure' else 'table_title')]
        if len(regions)!=1 or len(captions)!=1: raise ValueError(f'Ambiguous visual/caption on page {num}: {len(regions)}, {len(captions)}')
        page=pdf[num-1]; sx=page.rect.width/data['page_width']; sy=page.rect.height/data['page_height']
        def rect(v):
            x0,y0,x1,y1=v['bbox']; return fitz.Rect(x0*sx,y0*sy,x1*sx,y1*sy)
        crop=rect(regions[0]); caption=page.get_text('text',clip=rect(captions[0])).strip()
        # Captions are source-PDF text, not generated associations.
        if not re.search(r'(?:Fig\.|Table)\s*'+asset_id+r'\s*:',caption): raise ValueError(f'Caption ID mismatch: {caption[:80]}')
        dest=assets/f'{kind}-{asset_id}.png'; pix=page.get_pixmap(matrix=fitz.Matrix(4,4),clip=crop,alpha=False);pix.save(dest)
        record={'caption':caption,'path':str(dest),'width':pix.width,'height':pix.height,'aspect':pix.width/pix.height,
                'source_pdf_page':num,'source_region_id':regions[0]['region_id'],'source_caption_region_id':captions[0]['region_id'],
                'source_bbox':regions[0]['bbox'],'pdf_clip':list(crop),'caption_association':'single matching visual and caption on source page; pilot selection',
                'source_flags':regions[0].get('flags',[])}
        (figures if kind=='figure' else tables)[asset_id]=record
    (content/'raw.md').write_text('\n\n'.join(raw))
    write(assets/'figures.json',figures);write(assets/'tables.json',tables)
    write(out/'input_manifest.json',{'paper':str(pdf_path),'paper_sha256':sha(pdf_path),'source_pages':records,
          'scope':'main paper pages 1-17; references and appendix excluded; five explicitly selected source visuals',
          'text_source':'existing DeepSeekOCR2 aligned/enriched Markdown',
          'visual_source':'existing PP-DocLayout boxes re-rendered from original PDF at 4x for legibility',
          'upstream_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=VENDOR,text=True).strip(),
          'adaptations':['existing parser output replaces Marker','local Qwen3.5 replaces cloud model','no institution/conference logos supplied'],
          'production_modified':False})
    print(f'Prepared {len(raw)} pages, {len(figures)} figures, {len(tables)} tables',flush=True)

class LocalAgent:
    llm=None
    def __init__(self, system_msg, config, state=None, agent_name='unknown'):
        self.system_msg=system_msg;self.state=state;self.agent_name=agent_name;self.history=[]
    def reset(self): self.history=[]
    def step(self,message):
        from evidence_graph.llm_client import create_llm
        out=Path(self.state['output_dir']); cache=out/'calls';cache.mkdir(exist_ok=True)
        key=hashlib.sha256((self.system_msg+message).encode()).hexdigest()[:20]
        dest=cache/f'{self.agent_name}_{key}.json'
        if dest.exists():
            result=json.loads(dest.read_text());print('cached',self.agent_name,key,flush=True)
        else:
            if LocalAgent.llm is None:
                LocalAgent.llm=create_llm({'provider':'transformers','model':str(ROOT/'external/models/Qwen3.5-35B-A3B'),
                    'dtype':'bfloat16','device_map':'auto','require_cuda':True,'enable_thinking':False,'max_new_tokens':6500})
            llm=LocalAgent.llm; image_paths=[]; prompt=message
            try: parts=json.loads(message)
            except ValueError: parts=None
            if isinstance(parts,list):
                prompt='\n'.join(p['text'] for p in parts if p.get('type')=='text')
                for i,p in enumerate(parts):
                    if p.get('type')=='image_url':
                        path=cache/f'{key}_{i}.png';path.write_bytes(base64.b64decode(p['image_url']['url'].split(',',1)[1]));image_paths.append(str(path))
            start=time.time();print('generate',self.agent_name,'prompt_chars',len(prompt),flush=True)
            # This method also preserves raw non-JSON output for debugging.
            gen=llm.generate_raw_with_images(self.system_msg,prompt,image_paths,max_new_tokens=6500)
            result={'content':gen.raw,'input_tokens':len(llm.tokenizer.encode(prompt)),
                    'output_tokens':len(llm.tokenizer.encode(gen.raw)), 'seconds':time.time()-start,
                    'model':llm.model_name,'system':self.system_msg,'prompt':prompt,'vision_images':image_paths,
                    'token_accounting':'text token counts only; image and chat-template tokens excluded'}
            write(dest,result);print('done',self.agent_name,round(result['seconds'],1),flush=True)
        self.state['timing_metrics'].add_api_call(self.agent_name,'local',result['input_tokens'],result['output_tokens'])
        return types.SimpleNamespace(**{k:result[k] for k in ['content','input_tokens','output_tokens']})

def install_bridge():
    import json_repair
    module=types.ModuleType('utils.langgraph_utils')
    module.LangGraphAgent=LocalAgent
    module.extract_json=lambda text:json_repair.loads(text)
    module.load_prompt=lambda path:(VENDOR/path).read_text()
    sys.modules['utils.langgraph_utils']=module

def generate(out):
    install_bridge()
    from jinja2 import Template
    from src.state.poster_state import create_state,ModelConfig
    from utils.langgraph_utils import extract_json
    from src.agents.curator import curator_node
    from src.agents.color_agent import color_agent_node
    from src.agents.section_title_designer import section_title_designer_node
    from src.agents.layout_with_balancer import layout_with_balancer_node
    from src.agents.font_agent import font_agent_node
    from src.agents.renderer import renderer_node
    manifest=json.loads((out/'input_manifest.json').read_text())
    state=create_state(manifest['paper'],width=54,height=36,url=manifest.get('url',''))
    state.update(output_dir=str(out),poster_name=manifest.get('poster_name','detr_postergen_qwen35b'),
                 text_model=ModelConfig('Qwen3.5-35B-A3B','local',temperature=0,max_tokens=6500),
                 vision_model=ModelConfig('Qwen3.5-35B-A3B','local',temperature=0,max_tokens=6500))
    content=out/'content';raw=(content/'raw.md').read_text()
    state['raw_text']=raw
    state['images']=json.loads((out/'assets/figures.json').read_text())
    state['tables']=json.loads((out/'assets/tables.json').read_text())
    def ask(template,**kwargs):
        prompt=Template((VENDOR/'config/prompts'/template).read_text()).render(**kwargs)
        return extract_json(LocalAgent('expert academic paper parser',state['text_model'],state,'parser').step(prompt).content)
    meta=ask('extract_title_authors.txt',markdown_document=raw[:12000])
    narrative=ask('narrative_abt_extraction.txt',markdown_document=raw)
    narrative['meta']={'poster_title':meta['title'],'authors':meta['authors']}
    sections=ask('extract_structured_sections.txt',raw_text=raw)
    visuals=[{'id':f'{kind}_{k}','type':kind,'caption':v['caption'],'aspect_ratio':v['aspect']}
             for kind,group in [('figure',state['images']),('table',state['tables'])] for k,v in group.items()]
    classified=ask('classify_visuals.txt',visuals_list=json.dumps(visuals))
    for key,val in [('narrative_content',narrative),('structured_sections',sections),('classified_visuals',classified)]:
        state[key]=val;write(content/f'{key}.json',val)
    for node in [curator_node,color_agent_node,section_title_designer_node,layout_with_balancer_node,font_agent_node,renderer_node]:
        start=time.time();state=node(state);write(out/'state_checkpoint.json',state)
        if state.get('errors'): raise RuntimeError(state['errors'])
        print('STAGE',node.__name__,round(time.time()-start,1),flush=True)
    write(out/'run_status.json',{'status':'pptx_generated','model':'Qwen3.5-35B-A3B','local_calls':len(list((out/'calls').glob('*.json'))),
        'official_cloud_baseline':False,'limitations':['existing OCR and caption-association review flags remain','local-model adapted run, not official benchmark reproduction']})

def render(out):
    import fitz
    from xml.sax.saxutils import escape
    office=HERE/'libreoffice/opt/libreoffice25.8/program/soffice'
    manifest=json.loads((out/'input_manifest.json').read_text())
    pptx=out/(manifest.get('poster_name','detr_postergen_qwen35b')+'.pptx')
    # Use the same upstream fonts as layout measurement, without installing
    # fonts globally or writing to the user's font cache.
    (HERE/'font-cache').mkdir(exist_ok=True)
    fonts=out/'fonts.conf'
    fonts.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
                     '<include>/etc/fonts/fonts.conf</include><dir>'+escape(str(VENDOR/'fonts'))+'</dir>'
                     '<cachedir>'+escape(str(HERE/'font-cache'))+'</cachedir></fontconfig>')
    env=dict(os.environ,SAL_USE_VCLPLUGIN='svp',FONTCONFIG_FILE=str(fonts))
    profile='file:///tmp/postergen-lo-'+hashlib.sha256(str(out).encode()).hexdigest()[:16]
    subprocess.run([str(office),'-env:UserInstallation='+profile,'--headless','--convert-to','pdf','--outdir',str(out),str(pptx)],check=True,timeout=120,env=env)
    pdf=fitz.open(pptx.with_suffix('.pdf'));page=pdf[0]
    page.get_pixmap(matrix=fitz.Matrix(2400/page.rect.width,2400/page.rect.width),alpha=False).save(out/'preview.png')
    print('Rendered',out/'preview.png')

def audit(out):
    import fitz
    from pptx import Presentation
    manifest=json.loads((out/'input_manifest.json').read_text())
    pptx=out/(manifest.get('poster_name','detr_postergen_qwen35b')+'.pptx')
    prs=Presentation(pptx); pdf=fitz.open(pptx.with_suffix('.pdf'))
    text='\n'.join(p.get_text() for p in pdf)
    (out/'rendered_text.txt').write_text(text)
    def norm(s):
        import unicodedata
        return re.sub(r'[^\w]+','',unicodedata.normalize('NFKC',s)).casefold()
    missing=[]; bounds=[]; ntext=0;npictures=0
    for slide in prs.slides:
        for i,shape in enumerate(slide.shapes):
            if (shape.left < -9144 or shape.top < -9144 or
                shape.left+shape.width > prs.slide_width+9144 or
                shape.top+shape.height > prs.slide_height+9144):
                bounds.append({'shape':i,'name':shape.name})
            npictures+=int(shape.shape_type==13)
            if shape.has_text_frame and shape.text.strip():
                ntext+=1
                for p in shape.text_frame.paragraphs:
                    if len(norm(p.text))>3 and norm(p.text) not in norm(text):
                        missing.append({'shape':i,'text':p.text})
    # Render-level checks complement, but do not replace, source review.
    storyboard=json.loads((out/'content/optimized_story_board.json').read_text())
    sections=storyboard['spatial_content_plan']['sections']
    selected=[v['visual_id'] for s in sections for v in s.get('visual_assets',[])]
    known={kind+'_'+key for kind,file in [('figure','figures.json'),('table','tables.json')]
           for key in json.loads((out/'assets'/file).read_text())}
    headings={norm(s['section_title']) for s in sections}
    collisions=[]
    for page in pdf:
        lines=[{'text':''.join(s['text'] for s in line['spans']),'bbox':line['bbox']}
               for block in page.get_text('dict')['blocks'] if 'lines' in block for line in block['lines']]
        for heading in lines:
            if norm(heading['text']) not in headings: continue
            for other in lines:
                if other is heading or norm(other['text']) in headings: continue
                overlap=fitz.Rect(heading['bbox']) & fitz.Rect(other['bbox'])
                if not overlap.is_empty and overlap.height>1:
                    collisions.append({'heading':heading['text'],'other_text':other['text'],'line_box_overlap_points':round(overlap.height,2)})
    markup=re.findall(r'\\(?:\(|\)|\[|\]|frac|sum|mathrm|text|sigma|times)|\*[^*\n]{2,40}\*|</?color[^>]*>',text)
    write(out/'structural_audit.json',{'slides':len(prs.slides),'pdf_pages':len(pdf),
        'size_inches':[prs.slide_width/914400,prs.slide_height/914400],
        'text_shapes':ntext,'picture_shapes_including_qr':npictures,
        'out_of_slide_shapes':bounds,'paragraphs_not_recovered_in_rendered_pdf':missing,
        'selected_visual_ids':selected,'unknown_visual_ids':sorted(set(selected)-known),
        'heading_line_box_collision_candidates':collisions,
        'visible_markup_candidates':markup,
        'rendered_font_names':sorted({span['font'] for page in pdf for block in page.get_text('dict')['blocks'] if 'lines' in block for line in block['lines'] for span in line['spans']}),
        'scope':'Geometry and rendered text recovery only; no automatic factual or aesthetic quality score.'})
    print(json.dumps({'out_of_slide_shapes':len(bounds),'missing_paragraphs':len(missing),'text_shapes':ntext,'picture_shapes':npictures}))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('stage',choices=['prepare','generate','render','audit'])
    parser.add_argument('--output',type=Path,default=HERE/'output/detr');args=parser.parse_args()
    args.output=args.output.resolve();args.output.mkdir(parents=True,exist_ok=True)
    globals()[args.stage](args.output)
